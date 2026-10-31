from pathlib import Path
from math import cos, sin, pi

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)

W, H = 13.333, 7.5

GREEN = "1F4B3F"
TERRA = "B5651D"
CREAM = "F5F3EF"
SAND = "EDEAE3"
INK = "2B2A28"
WHITE = "FFFFFF"
MUTED = "6F746F"
PALE_GREEN = "DCE7E2"
PALE_TERRA = "F1DED0"

FONT = "Noto Sans"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def rect(slide, x, y, w, h, fill, radius=False, line=None, line_width=1):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def circle(slide, x, y, d, fill, line=None, line_width=1):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=INK, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
            margin=0.05, italic=False, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def rich_text(slide, parts, x, y, w, h, size=18, color=INK,
              align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for part in parts:
        run = p.add_run()
        run.text = part[0]
        run.font.name = FONT
        run.font.size = Pt(part[1] if len(part) > 1 and part[1] else size)
        run.font.bold = bool(part[2]) if len(part) > 2 else False
        run.font.color.rgb = rgb(part[3] if len(part) > 3 and part[3] else color)
    return box


def bullets(slide, items, x, y, w, h, size=17, color=INK, bullet_color=None,
            gap=8, level=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.level = level
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        p.text = "•  " + p.text
    return box


def add_picture_cover(slide, path, x, y, w, h):
    path = str(path)
    with Image.open(path) as im:
        iw, ih = im.size
    img_ratio = iw / ih
    frame_ratio = w / h
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if img_ratio > frame_ratio:
        visible = frame_ratio / img_ratio
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    else:
        visible = img_ratio / frame_ratio
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2
    return pic


def add_picture_contain(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                    Inches(pw), Inches(ph))


def add_title(slide, kicker, title, number=None, dark=False, title_size=27):
    fg = WHITE if dark else INK
    muted = SAND if dark else TERRA
    textbox(slide, kicker.upper(), 0.72, 0.42, 8.0, 0.28, 10, muted, True)
    textbox(slide, title, 0.72, 0.78, 11.8, 0.72, title_size, fg, True)
    rect(slide, 0.72, 1.54, 0.78, 0.05, TERRA)
    if number is not None:
        textbox(slide, f"{number:02d}", 12.0, 0.45, 0.62, 0.3, 10, muted, True, PP_ALIGN.RIGHT)


def footer(slide, n, dark=False):
    color = SAND if dark else MUTED
    textbox(slide, "FOSE Safar · Projet de stage", 0.72, 7.12, 4.4, 0.2, 8, color)
    textbox(slide, str(n), 12.0, 7.12, 0.6, 0.2, 8, color, True, PP_ALIGN.RIGHT)


def pill(slide, text, x, y, w, fill, fg=WHITE, size=11):
    rect(slide, x, y, w, 0.38, fill, True)
    textbox(slide, text, x + 0.08, y + 0.01, w - 0.16, 0.31, size, fg, True,
            PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def stat_card(slide, value, label, x, y, w, fill=WHITE, accent=GREEN):
    rect(slide, x, y, w, 1.3, fill, True)
    rect(slide, x, y, 0.08, 1.3, accent, True)
    textbox(slide, value, x + 0.28, y + 0.20, w - 0.45, 0.46, 26, accent, True)
    textbox(slide, label, x + 0.28, y + 0.77, w - 0.45, 0.3, 11, MUTED, True)


def card(slide, title, body, x, y, w, h, accent=GREEN, fill=WHITE, num=None):
    rect(slide, x, y, w, h, fill, True)
    if num is not None:
        circle(slide, x + 0.28, y + 0.28, 0.55, accent)
        textbox(slide, str(num), x + 0.28, y + 0.285, 0.55, 0.45, 16, WHITE, True,
                PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        title_x = x + 1.0
        title_w = w - 1.28
    else:
        rect(slide, x + 0.28, y + 0.28, 0.08, 0.68, accent, True)
        title_x = x + 0.55
        title_w = w - 0.83
    textbox(slide, title, title_x, y + 0.25, title_w, 0.45, 16, INK, True)
    textbox(slide, body, x + 0.3, y + 0.91, w - 0.6, h - 1.1, 12, MUTED)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

img = ROOT / "static" / "img" / "residence"
logo = Path("/tmp/emsi_logo.png")


# 1 — Couverture
slide = prs.slides.add_slide(blank)
add_picture_cover(slide, img / "piscine.jpeg", 0, 0, W, H)
rect(slide, 0, 0, 5.7, H, GREEN)
rect(slide, 5.7, 0, 0.12, H, TERRA)
pill(slide, "PROJET DE STAGE", 0.72, 0.66, 1.75, TERRA)
textbox(slide, "FOSE\nSAFAR", 0.72, 1.35, 4.3, 1.55, 38, WHITE, True)
textbox(slide, "Application web de gestion des\nréservations d’appartements", 0.72, 3.12, 4.3, 0.95, 20, WHITE, True)
textbox(slide, "Digitaliser · Sécuriser · Simplifier", 0.72, 4.28, 4.2, 0.34, 12, SAND, True)
rect(slide, 0.72, 4.88, 4.25, 0.02, TERRA)
textbox(slide, "Réalisé par", 0.72, 5.16, 1.2, 0.25, 9, SAND, True)
textbox(slide, "Iliane Firachine  ·  Rayane Firachine", 0.72, 5.48, 4.55, 0.34, 14, WHITE, True)
textbox(slide, "Encadrante et présidente de la FOSE : Mme Siham Lafkir", 0.72, 5.96, 4.55, 0.42, 10, SAND)
if logo.exists():
    rect(slide, 0.72, 6.58, 3.33, 0.55, WHITE, True)
    add_picture_contain(slide, logo, 0.87, 6.64, 3.03, 0.40)


# 2 — Le contexte
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Contexte", "Un patrimoine social à rendre accessible", 2)
add_picture_cover(slide, img / "balcon.jpeg", 0.72, 1.9, 5.05, 4.6)
rect(slide, 0.94, 5.63, 2.95, 0.58, GREEN, True)
textbox(slide, "Résidence Amwaj · Sidi Rahal", 1.08, 5.74, 2.65, 0.25, 11, WHITE, True,
        PP_ALIGN.CENTER)
stat_card(slide, "12", "APPARTEMENTS FOSE", 6.18, 1.94, 1.8, WHITE, GREEN)
stat_card(slide, "2", "PROFILS CLIENTS", 8.25, 1.94, 1.8, WHITE, TERRA)
stat_card(slide, "1", "PLATEFORME UNIQUE", 10.32, 1.94, 2.05, WHITE, GREEN)
textbox(slide, "La FOSE de l’enseignement Hay Hassani met en location ses appartements au sein d’une résidence privée, avec un avantage tarifaire destiné aux adhérents de la Fondation Mohammed VI.",
        6.18, 3.62, 6.15, 1.35, 16, INK, True)
bullets(slide, ["clients FM6 : tarif préférentiel et règles de quota",
                "clients réguliers : tarif standard, sans quota annuel",
                "administration : pilotage centralisé du parc"],
        6.18, 5.16, 5.95, 1.25, 13, MUTED, gap=6)
footer(slide, 2)


# 3 — Problématique
slide = prs.slides.add_slide(blank)
set_bg(slide, SAND)
add_title(slide, "Problématique", "Pourquoi digitaliser les réservations ?", 3)
textbox(slide, "AVANT", 0.82, 1.94, 1.1, 0.28, 11, TERRA, True)
problems = [
    ("Disponibilités", "Vérifications manuelles et risque de double réservation"),
    ("Tarification", "Calculs différents selon le profil du client"),
    ("Suivi", "Relances, paiements et justificatifs dispersés"),
    ("Conformité", "Collecte des données sans parcours formalisé"),
]
for i, (t, b) in enumerate(problems):
    y = 2.35 + i * 0.94
    circle(slide, 0.82, y, 0.52, TERRA)
    textbox(slide, "!", 0.82, y + 0.01, 0.52, 0.38, 18, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, t, 1.58, y - 0.02, 2.1, 0.28, 14, INK, True)
    textbox(slide, b, 1.58, y + 0.31, 3.75, 0.44, 11, MUTED)
line(slide, 5.82, 1.95, 5.82, 6.35, GREEN, 2)
circle(slide, 5.54, 3.75, 0.56, GREEN)
textbox(slide, "→", 5.54, 3.76, 0.56, 0.4, 20, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
textbox(slide, "APRÈS", 6.34, 1.94, 1.2, 0.28, 11, GREEN, True)
rect(slide, 6.3, 2.35, 6.05, 3.83, GREEN, True)
textbox(slide, "Un parcours numérique unique", 6.74, 2.78, 5.2, 0.46, 22, WHITE, True)
bullets(slide, ["disponibilité calculée en temps réel",
                "règles métier appliquées automatiquement",
                "traçabilité de chaque réservation",
                "e-mails, paiement et reçu intégrés",
                "tableau de bord pour l’administration"],
        6.74, 3.52, 5.1, 2.2, 15, WHITE, gap=8)
footer(slide, 3)


# 4 — Objectifs
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Vision", "Trois objectifs qui structurent la solution", 4)
cards = [
    ("01", "Simplifier", "Permettre au client de réserver en autonomie, depuis la recherche jusqu’au reçu.", GREEN),
    ("02", "Fiabiliser", "Bloquer les conflits de dates, automatiser les tarifs, quotas et expirations.", TERRA),
    ("03", "Piloter", "Donner à la FOSE une vision claire sur les logements, réservations et paiements.", GREEN),
]
for i, (n, t, b, c) in enumerate(cards):
    x = 0.78 + i * 4.18
    rect(slide, x, 2.0, 3.78, 3.92, WHITE, True)
    textbox(slide, n, x + 0.35, 2.35, 0.72, 0.45, 20, c, True)
    rect(slide, x + 0.35, 2.98, 0.62, 0.06, c)
    textbox(slide, t, x + 0.35, 3.36, 3.05, 0.46, 22, INK, True)
    textbox(slide, b, x + 0.35, 4.12, 3.05, 1.15, 15, MUTED)
    pill(slide, ["EXPÉRIENCE CLIENT", "RÈGLES MÉTIER", "DÉCISION"][i], x + 0.35, 5.26, 2.3, c, WHITE, 9)
textbox(slide, "La technologie sert le processus — et non l’inverse.", 0.82, 6.45, 11.7, 0.34, 15, GREEN, True, PP_ALIGN.CENTER)
footer(slide, 4)


# 5 — Acteurs
slide = prs.slides.add_slide(blank)
set_bg(slide, GREEN)
add_title(slide, "Utilisateurs", "Trois acteurs, des parcours clairement définis", 5, dark=True)
actor_data = [
    ("ADMIN", "Administrateur", ["gère les 12 appartements", "suit réservations et paiements", "consulte le tableau de bord"], TERRA),
    ("FM6", "Client FM6", ["réserve et suit ses séjours", "bénéficie du tarif réduit", "quota annuel contrôlé"], WHITE),
    ("STD", "Client régulier", ["réserve et paie en ligne", "reçoit confirmations et reçu", "aucun quota annuel"], SAND),
]
for i, (code, title, points, accent) in enumerate(actor_data):
    x = 0.78 + i * 4.18
    rect(slide, x, 2.03, 3.78, 4.35, WHITE, True)
    circle(slide, x + 0.36, 2.40, 0.78, accent if accent != WHITE else TERRA)
    textbox(slide, code, x + 0.36, 2.55, 0.78, 0.24, 10,
            GREEN if accent != TERRA else WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, title, x + 1.36, 2.45, 2.0, 0.42, 18, INK, True)
    rect(slide, x + 0.36, 3.36, 3.02, 0.02, SAND)
    bullets(slide, points, x + 0.36, 3.72, 3.02, 1.7, 13, MUTED, gap=11)
    if i == 1:
        pill(slide, "500 MAD / nuit", x + 0.36, 5.58, 1.66, TERRA, WHITE, 10)
    elif i == 2:
        pill(slide, "700 MAD / nuit", x + 0.36, 5.58, 1.66, GREEN, WHITE, 10)
footer(slide, 5, dark=True)


# 6 — Parcours réservation
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Parcours client", "Une réservation en quatre étapes", 6, title_size=24)
steps = [
    ("Choisir", "Dates + appartement", GREEN),
    ("Confirmer", "Récapitulatif + politiques par e-mail", TERRA),
    ("Payer", "Délai de 24 heures", GREEN),
    ("Recevoir", "Reçu PDF + pièces pour le check-in", TERRA),
]
for i, (t, b, c) in enumerate(steps):
    cx = 1.15 + i * 3.18
    if i < 3:
        line(slide, cx + 1.0, 3.04, cx + 2.24, 3.04, SAND, 6)
    circle(slide, cx, 2.22, 1.58, c)
    textbox(slide, str(i + 1), cx, 2.49, 1.58, 0.46, 26, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, t, cx - 0.36, 4.16, 2.3, 0.4, 17, INK, True, PP_ALIGN.CENTER)
    textbox(slide, b, cx - 0.43, 4.70, 2.45, 0.78, 12, MUTED, False, PP_ALIGN.CENTER)
rect(slide, 3.05, 5.90, 7.25, 0.64, PALE_TERRA, True)
textbox(slide, "Sans paiement sous 24 h : réservation expirée → appartement de nouveau disponible",
        3.28, 6.05, 6.8, 0.28, 12, TERRA, True, PP_ALIGN.CENTER)
footer(slide, 6)


# 7 — Règles métier
slide = prs.slides.add_slide(blank)
set_bg(slide, SAND)
add_title(slide, "Règles métier", "L’équité est encodée dans le système", 7)
rect(slide, 0.78, 1.94, 5.78, 4.7, WHITE, True)
pill(slide, "CLIENT FM6", 1.18, 2.30, 1.55, TERRA)
rich_text(slide, [("500", 32, True, TERRA), (" MAD / nuit", 16, True, INK)], 1.18, 3.03, 3.2, 0.6)
bullets(slide, ["5 nuits maximum par réservation",
                "10 nuits maximum par année civile",
                "matricule obligatoire à l’inscription",
                "contrôles automatiques avant validation"],
        1.18, 3.88, 4.78, 1.85, 14, MUTED, gap=9)
rect(slide, 6.78, 1.94, 5.78, 4.7, GREEN, True)
pill(slide, "CLIENT RÉGULIER", 7.18, 2.30, 1.95, TERRA)
rich_text(slide, [("700", 32, True, WHITE), (" MAD / nuit", 16, True, SAND)], 7.18, 3.03, 3.2, 0.6, color=WHITE)
bullets(slide, ["même parcours de réservation",
                "mêmes services et notifications",
                "aucun quota annuel",
                "tarification standard"],
        7.18, 3.88, 4.78, 1.85, 14, WHITE, gap=9)
textbox(slide, "Disponibilité, chevauchements et montant total sont recalculés côté serveur.",
        2.1, 6.75, 9.1, 0.25, 11, GREEN, True, PP_ALIGN.CENTER)
footer(slide, 7)


# 8 — Fonctionnalités
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Produit", "Les fonctionnalités clés de FOSE Safar", 8)
items = [
    ("Catalogue", "Recherche par dates, galerie et prix personnalisé", GREEN),
    ("Réservation", "Numéro unique, historique et suivi des statuts", TERRA),
    ("Paiement", "Carte, confirmation et reçu PDF envoyé par e-mail", GREEN),
    ("Administration", "Logements, périodes, réservations et paiements", TERRA),
    ("Notifications", "Politiques, confirmation, expiration et alertes", GREEN),
    ("Confidentialité", "Information, consentement horodaté et droits", TERRA),
]
for i, (t, b, c) in enumerate(items):
    col, row = i % 3, i // 3
    x, y = 0.78 + col * 4.18, 1.94 + row * 2.28
    card(slide, t, b, x, y, 3.78, 1.88, c, WHITE)
rect(slide, 1.8, 6.72, 9.75, 0.32, PALE_GREEN, True)
textbox(slide, "Une seule source de vérité pour le client et l’administrateur", 2.0, 6.76, 9.35, 0.2, 10, GREEN, True, PP_ALIGN.CENTER)
footer(slide, 8)


# 9 — Architecture
slide = prs.slides.add_slide(blank)
set_bg(slide, GREEN)
add_title(slide, "Architecture", "Une application Django modulaire et maintenable", 9, dark=True)
layers = [
    ("Navigateur", "Interfaces responsives", WHITE, GREEN),
    ("Django", "Vues · formulaires · authentification", TERRA, WHITE),
    ("Services métier", "Tarifs · quotas · expiration · e-mails", WHITE, GREEN),
    ("Données", "ORM · SQLite / PostgreSQL", TERRA, WHITE),
]
for i, (t, b, fill, fg) in enumerate(layers):
    x = 0.76 + i * 3.16
    rect(slide, x, 2.32, 2.62, 1.45, fill, True)
    textbox(slide, t, x + 0.18, 2.62, 2.26, 0.34, 16, fg, True, PP_ALIGN.CENTER)
    textbox(slide, b, x + 0.18, 3.06, 2.26, 0.42, 10, fg, False, PP_ALIGN.CENTER)
    if i < 3:
        textbox(slide, "→", x + 2.66, 2.78, 0.48, 0.42, 25, SAND, True, PP_ALIGN.CENTER)
apps = ["accounts", "apartments", "reservations", "payments", "dashboard", "notifications"]
for i, label in enumerate(apps):
    x = 0.82 + i * 2.05
    y = 4.74 + (i % 3) * 0.0
    pill(slide, label, x, y, 1.77, TERRA if i % 2 else WHITE, GREEN if i % 2 == 0 else WHITE, 10)
textbox(slide, "Séparation par domaines", 0.94, 4.25, 2.5, 0.25, 11, SAND, True)
textbox(slide, "Sécurité native", 0.94, 5.57, 2.2, 0.25, 11, SAND, True)
bullets(slide, ["CSRF et sessions Django", "mots de passe hachés", "contrôles d’accès par rôle", "transactions atomiques"],
        0.94, 5.91, 5.15, 0.98, 11, WHITE, gap=3)
textbox(slide, "Optimisation", 7.02, 5.57, 2.2, 0.25, 11, SAND, True)
bullets(slide, ["requêtes relationnelles optimisées", "validation centralisée", "pagination et filtres", "tests automatisés"],
        7.02, 5.91, 5.0, 0.98, 11, WHITE, gap=3)
footer(slide, 9, dark=True)


# 10 — Visuels résidence
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Expérience", "Un séjour valorisé par une présentation visuelle", 10)
photos = [
    (img / "salon-appartement.jpeg", 0.72, 1.88, 5.65, 2.18, "Salon"),
    (img / "chambre-king.jpeg", 6.56, 1.88, 2.83, 2.18, "Chambre"),
    (img / "cuisine-appartement.jpeg", 9.58, 1.88, 3.03, 2.18, "Cuisine"),
    (img / "piscine-jardin.jpeg", 0.72, 4.28, 2.83, 2.12, "Résidence"),
    (img / "terrain-sport.jpeg", 3.74, 4.28, 2.63, 2.12, "Équipements"),
    (img / "piscine-nuit.jpeg", 6.56, 4.28, 6.05, 2.12, "Piscine"),
]
for pth, x, y, w, h, label in photos:
    add_picture_cover(slide, pth, x, y, w, h)
    rect(slide, x + 0.12, y + h - 0.48, 1.1, 0.32, GREEN, True)
    textbox(slide, label, x + 0.17, y + h - 0.43, 1.0, 0.2, 9, WHITE, True, PP_ALIGN.CENTER)
footer(slide, 10)


# 11 — Sécurité
slide = prs.slides.add_slide(blank)
set_bg(slide, SAND)
add_title(slide, "Qualité", "Sécurité et fiabilité dès la conception", 11)
security = [
    ("Authentification", "Compte e-mail, sessions et mots de passe robustes", GREEN),
    ("Autorisation", "Accès séparés selon le rôle et les objets du client", TERRA),
    ("Intégrité", "Transactions atomiques et contrôles de cohérence", GREEN),
    ("Protection web", "CSRF, validation serveur et en-têtes Django", TERRA),
]
for i, (t, b, c) in enumerate(security):
    x = 0.78 + (i % 2) * 6.03
    y = 1.94 + (i // 2) * 1.72
    card(slide, t, b, x, y, 5.66, 1.38, c, WHITE)
rect(slide, 0.78, 5.52, 11.69, 1.05, GREEN, True)
textbox(slide, "Principe", 1.14, 5.74, 1.2, 0.26, 10, SAND, True)
textbox(slide, "Ne jamais faire confiance au navigateur : toutes les règles critiques sont revérifiées côté serveur.",
        2.45, 5.72, 9.55, 0.44, 16, WHITE, True)
footer(slide, 11)


# 12 — Conformité
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Conformité", "La loi 09-08 a influencé le produit", 12)
stages = [
    ("Cartographier", "Quelles données ? Pour quelles finalités ?", GREEN),
    ("Notifier", "Déclaration CNDP — formulaire F214 selon le régime applicable", TERRA),
    ("Informer", "Politique de confidentialité claire et accessible", GREEN),
    ("Permettre", "Accès, rectification, opposition et suppression", TERRA),
]
for i, (t, b, c) in enumerate(stages):
    x = 0.84 + i * 3.1
    circle(slide, x + 0.76, 2.25, 0.86, c)
    textbox(slide, str(i + 1), x + 0.76, 2.42, 0.86, 0.3, 18, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 3:
        line(slide, x + 1.63, 2.68, x + 3.12, 2.68, SAND, 5)
    textbox(slide, t, x, 3.38, 2.4, 0.38, 16, INK, True, PP_ALIGN.CENTER)
    textbox(slide, b, x, 3.88, 2.4, 0.82, 11, MUTED, False, PP_ALIGN.CENTER)
rect(slide, 0.88, 5.16, 11.58, 1.14, PALE_TERRA, True)
textbox(slide, "Hébergement hors du Maroc ?", 1.18, 5.42, 2.75, 0.32, 14, TERRA, True)
textbox(slide, "Évaluer et notifier le transfert de données à l’étranger via le formulaire F118 avant mise en production.",
        4.08, 5.38, 7.92, 0.48, 13, INK, True)
textbox(slide, "Source : CNDP — cndp.ma (consulté le 14/09/2026)", 0.92, 6.63, 5.0, 0.2, 8, MUTED)
footer(slide, 12)


# 13 — Démarche ingénieur
slide = prs.slides.add_slide(blank)
set_bg(slide, GREEN)
add_title(slide, "Apprentissage", "Un stage au-delà de l’exécution technique", 13, dark=True)
textbox(slide, "Nous n’avons pas seulement codé une demande.", 0.82, 1.95, 11.7, 0.52, 24, WHITE, True, PP_ALIGN.CENTER)
pillars = [
    ("Comprendre", "Observer le processus réel et ses contraintes", TERRA),
    ("Proposer", "Comparer les options et défendre des choix", WHITE),
    ("Concevoir", "Transformer les règles métier en contrôles fiables", TERRA),
    ("Responsabiliser", "Intégrer sécurité, données et conformité", WHITE),
]
for i, (t, b, c) in enumerate(pillars):
    x = 0.78 + i * 3.15
    if i < 3:
        line(slide, x + 2.62, 3.26, x + 3.14, 3.26, SAND, 4)
        textbox(slide, "→", x + 2.62, 3.02, 0.52, 0.36, 18, SAND, True, PP_ALIGN.CENTER)
    circle(slide, x + 0.88, 2.72, 1.08, c if c != WHITE else SAND)
    textbox(slide, str(i + 1), x + 0.88, 2.94, 1.08, 0.34, 20,
            GREEN if c != TERRA else WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, t, x + 0.03, 4.03, 2.78, 0.36, 16, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, b, x + 0.03, 4.49, 2.78, 0.72, 11, SAND, False, PP_ALIGN.CENTER)
rect(slide, 3.96, 5.65, 5.42, 0.72, TERRA, True)
textbox(slide, "POSTURE D’INGÉNIEUR", 4.18, 5.83, 4.98, 0.28, 16, WHITE, True, PP_ALIGN.CENTER)
footer(slide, 13, dark=True)


# 14 — État & feuille de route
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_title(slide, "Feuille de route", "Du démonstrateur au service en production", 14)
columns = [
    ("ACQUIS", GREEN, ["parcours de réservation", "tarification FM6 / standard", "expiration à 24 h", "e-mails et reçu PDF", "back-office et tests"]),
    ("À FINALISER", TERRA, ["passerelle de paiement réelle", "validation d’éligibilité FM6", "dossier CNDP complet", "recette avec utilisateurs"]),
    ("MISE EN PRODUCTION", GREEN, ["nom de domaine", "hébergement choisi", "certificat SSL / HTTPS", "PostgreSQL et sauvegardes", "supervision et support"]),
]
for i, (title, c, pts) in enumerate(columns):
    x = 0.78 + i * 4.18
    rect(slide, x, 1.95, 3.78, 4.84, WHITE, True)
    rect(slide, x, 1.95, 3.78, 0.66, c, True)
    textbox(slide, title, x + 0.2, 2.13, 3.38, 0.25, 12, WHITE, True, PP_ALIGN.CENTER)
    bullets(slide, pts, x + 0.35, 2.98, 3.08, 2.7, 13, MUTED, gap=10)
    textbox(slide, ["✓", "…", "→"][i], x + 1.45, 5.92, 0.88, 0.56, 30, c, True, PP_ALIGN.CENTER)
footer(slide, 14)


# 15 — Conclusion
slide = prs.slides.add_slide(blank)
add_picture_cover(slide, img / "piscine-nuit.jpeg", 0, 0, W, H)
rect(slide, 0, 0, W, H, GREEN)
rect(slide, 0.72, 0.72, 0.12, 5.95, TERRA)
textbox(slide, "CONCLUSION", 1.2, 0.92, 2.2, 0.3, 11, TERRA, True)
textbox(slide, "Ce projet n’était pas\nseulement technique.", 1.2, 1.48, 10.8, 1.25, 34, WHITE, True)
textbox(slide, "Il nous a appris à relier le besoin métier, l’expérience utilisateur, la sécurité, la conformité juridique et la réalité d’un déploiement.",
        1.2, 3.02, 10.25, 1.05, 20, SAND, True)
rich_text(slide, [("Une application utile naît d’un code solide — ", 18, False, WHITE),
                  ("et de décisions responsables.", 18, True, TERRA)],
          1.2, 4.48, 10.6, 0.56, 18, WHITE)
pill(slide, "MERCI POUR VOTRE ATTENTION", 1.2, 5.53, 3.55, TERRA, WHITE, 11)
textbox(slide, "Iliane Firachine  ·  Rayane Firachine", 1.2, 6.22, 5.0, 0.32, 12, WHITE, True)


# 16 — Sources
slide = prs.slides.add_slide(blank)
set_bg(slide, SAND)
add_title(slide, "Annexe", "Sources et périmètre de la présentation", 16)
sources = [
    ("Projet", "Code source Django, migrations, tests et rapport technique du dépôt FOSE Safar."),
    ("Photographies", "Visuels de la résidence fournis avec le projet dans static/img/residence/."),
    ("Conformité", "Commission Nationale de Contrôle de la Protection des Données à Caractère Personnel — cndp.ma."),
    ("Logo", "École Marocaine des Sciences de l’Ingénieur — emsi.ma / portail étudiant EMSI."),
]
for i, (t, b) in enumerate(sources):
    y = 1.95 + i * 1.15
    pill(slide, f"{i+1:02d}", 0.84, y, 0.58, GREEN if i % 2 == 0 else TERRA, WHITE, 10)
    textbox(slide, t, 1.68, y - 0.01, 1.55, 0.32, 14, INK, True)
    textbox(slide, b, 3.32, y - 0.02, 8.75, 0.58, 12, MUTED)
rect(slide, 0.84, 6.52, 11.55, 0.38, PALE_GREEN, True)
textbox(slide, "Les tarifs et règles présentés correspondent à l’état du code au 14 septembre 2026.",
        1.02, 6.59, 11.15, 0.2, 9, GREEN, True, PP_ALIGN.CENTER)
footer(slide, 16)


# Document properties
prs.core_properties.title = "FOSE Safar — Présentation du projet de stage"
prs.core_properties.subject = "Application web de gestion des réservations des appartements FOSE à la résidence Amwaj, Sidi Rahal"
prs.core_properties.author = "Iliane Firachine et Rayane Firachine"
prs.core_properties.keywords = "FOSE, EMSI, Django, réservation, appartements, FM6, loi 09-08"

dest = OUT / "presentation_projet_stage_fose_safar.pptx"
prs.save(dest)
print(dest)
